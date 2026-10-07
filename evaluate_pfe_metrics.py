#!/usr/bin/env python3
"""
evaluate_pfe_metrics.py
=======================
Computes, automatically, the evaluation metrics of BOTH manuscripts from the CSV
written by run_unified_benchmarks.py (results_unified.csv):

  PFE 1 - Multi-Objective RL (MORL)                (manuscript section 5.5)
      Compiler metrics : execution time (s), rotation key size (MB)
      MORL metrics     : Cardinality IC, Hypervolume HV (normalised by 1.21),
                         Sparsity SP (normalised by 2), Expected Utility EU

  PFE 2 - Constrained RL + safety rollback          (manuscript section 5.3.3)
      Cost Reduction CR%, Violation count / rate, Safe CR%, Safe violation count,
      within-budget vs above-budget breakdown, and the noise-estimator accuracy
      (MAE, MSE, RMSE, R2 - manuscript Table 5.3)

Usage
-----
    python3 evaluate_pfe_metrics.py --csv results_unified.csv
    python3 evaluate_pfe_metrics.py --csv results_unified.csv --bounds my_bounds.csv
    python3 evaluate_pfe_metrics.py --selftest

Output: eval_out/evaluation_metrics.xlsx  (sheets: README, MORL_metrics, Compiler_metrics,
Constrained_metrics, Noise_estimator, MORL_points, Raw_data, Plots)
"""
import argparse
import os
import sys

import numpy as np
import pandas as pd

Z_REF = 1.1                      # Nadir reference point (manuscript eq. 5.2)
HV_MAX = Z_REF * Z_REF           # 1.21
SP_MAX = 2.0                     # theoretical max squared distance (eq. 5.6)

# Table 5.2 of the MORL manuscript: (Tmin, Tmax, Kmax). Names follow the benchmark
# folders of the repo (dot_product, l2_distance, hamming_dist, lin_reg, ...).
# Poly. Tree benchmarks are not listed (no fixed name in the repo).
TABLE_5_2 = {
    "dot_product_4": (215, 1756, 3), "dot_product_8": (284, 3757, 4),
    "dot_product_16": (386, 7758, 5), "dot_product_32": (503, 15759, 5),
    "l2_distance_4": (218, 3757, 3), "l2_distance_8": (287, 7758, 4),
    "l2_distance_16": (391, 15759, 5), "l2_distance_32": (598, 31760, 5),
    "hamming_dist_4": (453, 4759, 3), "hamming_dist_8": (618, 9760, 4),
    "hamming_dist_16": (756, 19761, 5), "hamming_dist_32": (1065, 39762, 5),
    "lin_reg_4": (108, 3006, 2), "lin_reg_8": (108, 6006, 2),
    "lin_reg_16": (108, 12006, 1), "lin_reg_32": (108, 24006, 0),
    "poly_reg_4": (212, 5009, 3), "poly_reg_8": (212, 10009, 2),
    "poly_reg_16": (212, 20009, 1), "poly_reg_32": (211, 40009, 0),
    "matrix_mul_4": (275, 28006, 1), "box_blur_4": (14, 21006, 1),
    "gx_kernel_4": (381, 21507, 1), "gy_kernel_4": (381, 21507, 1),
    "roberts_cross_4": (181, 31007, 1),
    "max_3": (386, 3009, 4), "max_4": (793, 7012, 5), "max_5": (1619, 15015, 4),
    "sort_3": (912, 5012, 5),
}
ALIASES = {"hamming_distance": "hamming_dist", "linear_regression": "lin_reg",
           "polynomial_regression": "poly_reg", "matrix_multiplication": "matrix_mul",
           "sobel_gx": "gx_kernel", "sobel_gy": "gy_kernel", "maximum": "max"}


# --------------------------------------------------------------------------- #
# Pure metric functions (MORL)
# --------------------------------------------------------------------------- #
def pareto_front(points):
    """Non-dominated, de-duplicated subset (both objectives minimised), sorted by 1st obj."""
    pts = sorted({(float(a), float(b)) for a, b in points})
    front, best_y = [], np.inf
    for x, y in pts:                      # ascending x: keep only strictly improving y
        if y < best_y:
            front.append((x, y))
            best_y = y
    return front


def hypervolume(front, z_ref=Z_REF):
    """Eq. 5.3 on normalised points, divided by z_ref^2 (eq. 5.4)."""
    pts = [(x, y) for x, y in sorted(front) if x < z_ref and y < z_ref]
    hv = 0.0
    for i, (x, y) in enumerate(pts):
        nxt = pts[i + 1][0] if i + 1 < len(pts) else z_ref
        hv += (nxt - x) * (z_ref - y)
    return hv / (z_ref * z_ref)


def sparsity(front):
    """Eq. 5.5 normalised by 2 (eq. 5.6). 0 for a single point."""
    pts = sorted(front)
    if len(pts) < 2:
        return 0.0
    sq = sum((b[0] - a[0]) ** 2 + (b[1] - a[1]) ** 2 for a, b in zip(pts[:-1], pts[1:]))
    return sq / (len(pts) - 1) / SP_MAX


def expected_utility(front, grid=2001):
    """Eq. 5.7: E over w_exec~U[0,1] (w_keys = 1-w_exec) of min_x w.(T^,K^)."""
    w = np.linspace(0.0, 1.0, grid)
    arr = np.array(front)
    util = np.min(w[:, None] * arr[None, :, 0] + (1 - w[:, None]) * arr[None, :, 1], axis=1)
    return float(np.sum((util[1:] + util[:-1]) / 2 * np.diff(w)))       # trapezoid on [0,1]


def morl_metrics_from_normalised(norm_points):
    front = pareto_front(norm_points)
    return {"IC": len(front), "HV": hypervolume(front), "SP": sparsity(front),
            "EU": expected_utility(front)}, front


def canon(name):
    n = str(name).strip().lower().replace(" ", "_").replace("-", "_")
    for a, b in ALIASES.items():
        if n.startswith(a):
            n = b + n[len(a):]
    return n


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def load(csv_path, initial_budget):
    if str(csv_path).lower().endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(csv_path)
    else:
        df = pd.read_csv(csv_path)
    df = df.replace("N/A", np.nan)
    for c in df.columns:
        if c != "benchmark":
            df[c] = pd.to_numeric(df[c], errors="coerce")
    if "Remaining_noise_budget" in df:           # recompute so --initial-budget is honoured
        df["noise_used_measured"] = initial_budget - df["Remaining_noise_budget"]
    df["bench_key"] = df["benchmark"].map(canon)
    return df


def load_bounds(path):
    b = pd.read_excel(path) if str(path).lower().endswith((".xlsx", ".xlsm")) else pd.read_csv(path)
    b["bench_key"] = b["benchmark"].map(canon)
    return {r.bench_key: (float(r.Tmin), float(r.Tmax), float(r.Kmax)) for r in b.itertuples()}


# --------------------------------------------------------------------------- #
# PFE 1 : MORL
# --------------------------------------------------------------------------- #
def morl_report(df, t_col, k_col, user_bounds, eu_grid):
    rows, point_rows, compiler_rows, warnings = [], [], [], []
    if k_col in df and df[k_col].fillna(0).eq(0).all() and df.get("rotation_keys_static", pd.Series([0])).max() > 0:
        warnings.append(f"`{k_col}` is 0 for every run although the circuits use rotation keys: "
                        "the w_keys term is probably missing in pytrs/cost.py (calculate_cost). "
                        "Key objective will look constant; MORL metrics are not meaningful until fixed.")
    for bench, gb in df.groupby("benchmark"):
        key = canon(bench)
        sub_all = gb.dropna(subset=[t_col, k_col])
        if sub_all.empty:
            continue
        if key in user_bounds:
            (tmin, tmax, kmax), src = user_bounds[key], "bounds file"
        elif key in TABLE_5_2:
            (tmin, tmax, kmax), src = TABLE_5_2[key], "manuscript Table 5.2"
        else:
            (tmin, tmax, kmax), src = (sub_all[t_col].min(), sub_all[t_col].max(), sub_all[k_col].max()), "empirical (this CSV)"
            warnings.append(f"{bench}: no Table 5.2 bounds, using empirical bounds from this CSV "
                            "(HV/EU are then relative to your own runs only).")
        for budget, g in sub_all.groupby(sub_all["noise_budget"].fillna(-1)):
            t_hat = (g[t_col] - tmin) / (tmax - tmin) if tmax > tmin else g[t_col] * 0.0
            k_hat = g[k_col] / kmax if kmax > 0 else g[k_col] * 0.0
            clipped = bool(((t_hat < 0) | (t_hat > 1) | (k_hat > 1)).any())
            t_hat, k_hat = t_hat.clip(0, 1), k_hat.clip(0, 1)
            m, front = morl_metrics_from_normalised(list(zip(t_hat, k_hat)))
            fset = set(front)
            for (_, r), th, kh in zip(g.iterrows(), t_hat, k_hat):
                point_rows.append({"benchmark": bench, "noise_budget": budget, "w_ops": r.get("w_ops"),
                                   "w_keys": r.get("w_keys"), "T": r[t_col], "K": r[k_col],
                                   "T_hat": th, "K_hat": kh, "on_pareto_front": (th, kh) in fset})
            rows.append({"benchmark": bench, "noise_budget": budget, "n_runs": len(g),
                         "IC": m["IC"], "HV": round(m["HV"], 4), "SP": round(m["SP"], 4),
                         "EU": round(m["EU"], 4), "Tmin": tmin, "Tmax": tmax, "Kmax": kmax,
                         "bounds_source": src, "values_clipped": clipped})
            # compiler metrics at the two extreme preferences
            if {"w_ops", "w_keys"} <= set(g.columns):
                for label, order in (("speed (max w_ops)", ("w_ops", False)), ("memory (max w_keys)", ("w_keys", False))):
                    pick = g.sort_values(order[0], ascending=order[1]).iloc[0]
                    compiler_rows.append({
                        "benchmark": bench, "noise_budget": budget, "preference": label,
                        "w_ops": pick["w_ops"], "w_keys": pick["w_keys"],
                        "execution_time_s": pick.get("circuit_execution_time (s)"),
                        "total_time_s": pick.get("total_execution_time (s)"),
                        "rotation_key_size_MB": pick.get("rotation_keys_size (MB)"),
                        "rotation_keys": pick.get("rotation_keys_static")})
    return (pd.DataFrame(rows), pd.DataFrame(point_rows), pd.DataFrame(compiler_rows), warnings)


# --------------------------------------------------------------------------- #
# PFE 2 : constrained RL + safety rollback
# --------------------------------------------------------------------------- #
def cr(c0, cf):
    return (c0 - cf) / c0 * 100.0


def constrained_report(df, w_keys_filter, agent_name):
    warnings = []
    d = df.copy()
    if w_keys_filter is not None and "w_keys" in d and (d["w_keys"] == w_keys_filter).any():
        d = d[d["w_keys"] == w_keys_filter]
    elif w_keys_filter is not None:
        warnings.append(f"no run with w_keys={w_keys_filter}; constrained metrics use all runs.")
    d = d.dropna(subset=["noise_budget"])
    have_raw = {"initial_ops_cost", "agent_final_ops_cost", "agent_final_noise_estimated"} <= set(d.columns) \
        and d["initial_ops_cost"].notna().any()
    if not have_raw:
        warnings.append("columns initial_ops_cost / agent_final_* are missing: CR% and raw violations are "
                        "N/A. Apply run_py_extra_metrics.patch to RL/fhe_rl/run.py and re-run "
                        "run_unified_benchmarks.py. Only post-rollback (Safe) metrics are reported.")
    rows = []
    for budget, g in d.groupby("noise_budget"):
        has_init = "initial_noise_estimated" in g and g["initial_noise_estimated"].notna().any()
        within = g[g["initial_noise_estimated"] <= budget] if has_init else g.iloc[0:0]
        above = g[g["initial_noise_estimated"] > budget] if has_init else g.iloc[0:0]

        def block(sub, tag):
            out = {}
            n = len(sub)
            out[f"n{tag}"] = n
            if n == 0:
                return out
            safe_viol = int((sub["noise_used_estimated"] > budget).sum())
            out[f"safe_violations{tag}"] = safe_viol
            out[f"safe_viol_rate_%{tag}"] = round(100 * safe_viol / n, 1)
            if "initial_ops_cost" in sub:
                s2 = sub.dropna(subset=["initial_ops_cost", "final_ops_cost"])
                if len(s2):
                    out[f"safe_CR_%{tag}"] = round(cr(s2["initial_ops_cost"], s2["final_ops_cost"]).mean(), 2)
            if have_raw:
                s2 = sub.dropna(subset=["initial_ops_cost", "agent_final_ops_cost", "agent_final_noise_estimated"])
                if len(s2):
                    v = int((s2["agent_final_noise_estimated"] > budget).sum())
                    out[f"violations{tag}"] = v
                    out[f"viol_rate_%{tag}"] = round(100 * v / len(s2), 1)
                    out[f"CR_%{tag}"] = round(cr(s2["initial_ops_cost"], s2["agent_final_ops_cost"]).mean(), 2)
                    out[f"rollback_rate_%{tag}"] = round(100 * s2["rollback_applied"].mean(), 1)
            if "noise_used_measured" in sub:
                s2 = sub.dropna(subset=["noise_used_measured"])
                out[f"measured_violations{tag}"] = int((s2["noise_used_measured"] > budget).sum())
            return out

        row = {"agent": agent_name, "noise_budget": int(budget)}
        row.update(block(g, ""))
        if has_init:
            row["within_budget"] = len(within)
            row["above_budget"] = len(above)
            row.update({k: v for k, v in block(within, "_within").items() if k != "n_within"})
        rows.append(row)
    return pd.DataFrame(rows), warnings


def noise_estimator_report(df):
    d = df.dropna(subset=["noise_used_estimated", "noise_used_measured"])
    if len(d) == 0:
        return pd.DataFrame()
    y, yh = d["noise_used_measured"].to_numpy(float), d["noise_used_estimated"].to_numpy(float)
    err = yh - y
    ss_res, ss_tot = float(np.sum(err ** 2)), float(np.sum((y - y.mean()) ** 2))
    return pd.DataFrame([{"n": int(len(d)), "MAE": round(float(np.mean(np.abs(err))), 3),
                          "MSE": round(float(np.mean(err ** 2)), 3),
                          "RMSE": round(float(np.sqrt(np.mean(err ** 2))), 3),
                          "R2": round(1 - ss_res / ss_tot, 4) if ss_tot > 0 else float("nan"),
                          "mean_signed_error(est-measured)": round(float(err.mean()), 3)}])


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
def md_table(df):
    if df is None or df.empty:
        return "_no data_\n"
    cols = list(df.columns)
    lines = ["| " + " | ".join(map(str, cols)) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join("" if pd.isna(v) else str(v) for v in r.tolist()) + " |")
    return "\n".join(lines) + "\n"


def make_plots(points, df, out_dir):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        print("matplotlib not available: plots skipped")
        return
    os.makedirs(out_dir, exist_ok=True)
    for (bench, budget), g in points.groupby(["benchmark", "noise_budget"]):
        fig, ax = plt.subplots(figsize=(4.5, 4))
        dom, nd = g[~g.on_pareto_front], g[g.on_pareto_front].sort_values("T_hat")
        ax.scatter(dom.T_hat, dom.K_hat, c="gray", label="dominated")
        ax.scatter(nd.T_hat, nd.K_hat, c="red", label="non-dominated")
        ax.step(nd.T_hat, nd.K_hat, where="post", c="red", ls="--", lw=1)
        ax.set_xlabel("normalised execution cost T^"); ax.set_ylabel("normalised keys K^")
        ax.set_title(f"{bench} (budget {budget:g})", fontsize=9); ax.legend(fontsize=7)
        fig.tight_layout(); fig.savefig(os.path.join(out_dir, f"pareto_{bench}_{budget:g}.png"), dpi=130); plt.close(fig)
    d = df.dropna(subset=["noise_used_estimated", "noise_used_measured"])
    if len(d):
        fig, ax = plt.subplots(figsize=(4.5, 4))
        ax.scatter(d.noise_used_measured, d.noise_used_estimated, s=14)
        lim = [0, max(d.noise_used_measured.max(), d.noise_used_estimated.max()) * 1.1 + 1]
        ax.plot(lim, lim, "k--", lw=1); ax.set_xlabel("measured noise used (bits)")
        ax.set_ylabel("estimated noise used (bits)"); ax.set_title("Noise estimator accuracy", fontsize=9)
        fig.tight_layout(); fig.savefig(os.path.join(out_dir, "noise_estimator.png"), dpi=130); plt.close(fig)


README_METRICS = [
    ("IC (Cardinality)", "MORL", "Number of non-dominated (Pareto) circuits found, eq. 5.1", "higher"),
    ("HV (Hypervolume)", "MORL", "Dominated area, z_ref=(1.1,1.1), divided by 1.21, eq. 5.2-5.4", "higher (max 1)"),
    ("SP (Sparsity)", "MORL", "Mean squared gap between consecutive Pareto points, / 2, eq. 5.5-5.6", "lower (min 0)"),
    ("EU (Expected Utility)", "MORL", "Mean over w_exec~U[0,1] of the best weighted cost on the front, eq. 5.7", "lower (min 0)"),
    ("execution_time_s / rotation_key_size_MB", "MORL", "Measured by ./main (SEAL) at the speed / memory extreme preference", "lower"),
    ("CR_% / safe_CR_%", "Constrained", "(C0-Cf)/C0*100 for the agent output / after safety rollback", "higher"),
    ("violations / viol_rate_%", "Constrained", "Agent final estimated noise > budget (before rollback)", "lower"),
    ("safe_violations / safe_viol_rate_%", "Constrained", "Estimated noise > budget after safety rollback", "lower (0)"),
    ("within_budget / above_budget", "Constrained", "Runs whose INITIAL estimated noise is <= / > the budget", "-"),
    ("measured_violations", "Extension", "Runs where the real SEAL noise (initial budget - Remaining) exceeds the budget", "lower (0)"),
    ("MAE / MSE / RMSE / R2", "Noise estimator", "Estimated vs measured noise used (bits)", "MAE,MSE,RMSE lower; R2 higher"),
]


def write_workbook(path, sheets, notes, raw, points, df):
    """One Excel workbook: README + one sheet per table (+ embedded plots)."""
    import tempfile
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    tmp = tempfile.mkdtemp()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="README")
        for name, t in sheets.items():
            (t if len(t) else pd.DataFrame({"info": ["no data"]})).to_excel(xw, sheet_name=name, index=False)
        raw.drop(columns=["bench_key"], errors="ignore").to_excel(xw, sheet_name="Raw_data", index=False)
    wb = load_workbook(path)

    # README (built by hand)
    ws = wb["README"]
    ws.delete_rows(1, ws.max_row)
    rows = [("Evaluation metrics - MORL PFE and Constrained-RL PFE", None, None, None), ("", None, None, None),
            ("Metric", "PFE", "Meaning (manuscript equation)", "Better")]
    rows += README_METRICS
    rows += [("", None, None, None), ("Assumptions / notes", None, None, None)] + [(n, None, None, None) for n in notes]
    rows += [("", None, None, None), ("Sheets: MORL_metrics, Compiler_metrics, Constrained_metrics, Noise_estimator, "
                                      "MORL_points (normalised points + front flag), Raw_data (input), Plots.", None, None, None),
             ("Values are computed by evaluate_pfe_metrics.py from Raw_data (not live Excel formulas).", None, None, None)]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    for c in ws[3]:
        c.font, c.fill = bold, fill
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if c.font != bold:
                c.font = font
    ws.cell(row=len(README_METRICS) + 5, column=1).font = Font(name="Arial", size=10, bold=True)
    for col, w in zip("ABCD", (38, 16, 80, 26)):
        ws.column_dimensions[col].width = w

    # data sheets: header style, widths, freeze panes, Arial
    for name in list(sheets) + ["Raw_data"]:
        w = wb[name]
        w.freeze_panes = "A2"
        for c in w[1]:
            c.font, c.fill = bold, fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.font = font
        for i, col in enumerate(w.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            w.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 42)
        w.row_dimensions[1].height = 32

    # plots sheet
    if len(points):
        try:
            from openpyxl.drawing.image import Image
            make_plots(points, df, tmp)
            wp = wb.create_sheet("Plots")
            wp["A1"] = "Pareto fronts (normalised) and noise estimator accuracy"
            wp["A1"].font = Font(name="Arial", size=12, bold=True)
            files = sorted(f for f in os.listdir(tmp) if f.endswith(".png"))
            for i, fn in enumerate(files):
                img = Image(os.path.join(tmp, fn))
                img.width, img.height = 360, 320
                wp.add_image(img, f"{'ABCDEFGH'[(i % 3) * 3]}{3 + (i // 3) * 18}")
            wb.save(path)
        except Exception as exc:                       # plots are optional
            print("plots not embedded:", exc)
            wb.save(path)
    else:
        wb.save(path)
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", "--input", dest="csv", default="results_unified.csv",
                    help="benchmark results, .csv or .xlsx")
    ap.add_argument("--out", default="eval_out")
    ap.add_argument("--bounds", help="CSV with columns benchmark,Tmin,Tmax,Kmax (overrides Table 5.2)")
    ap.add_argument("--t-col", default="final_ops_cost", help="execution objective (default: optimizer cost)")
    ap.add_argument("--k-col", default="final_keys_cost", help="keys objective (default: optimizer keys cost)")
    ap.add_argument("--constrained-w-keys", type=float, default=0.0,
                    help="constrained-agent metrics use runs with this w_keys (speed focus); "
                         "use -1 for all runs")
    ap.add_argument("--agent-name", default="unified")
    ap.add_argument("--initial-budget", type=float, default=369.0,
                    help="SEAL initial noise budget assumed for noise_used_measured")
    ap.add_argument("--eu-grid", type=int, default=2001)
    ap.add_argument("--no-plots", action="store_true")
    ap.add_argument("--xlsx-name", default="evaluation_metrics.xlsx")
    ap.add_argument("--md", action="store_true", help="also write report.md")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()

    df = load(a.csv, a.initial_budget)
    bounds = load_bounds(a.bounds) if a.bounds else {}
    os.makedirs(a.out, exist_ok=True)

    morl, points, compiler, w1 = morl_report(df, a.t_col, a.k_col, bounds, a.eu_grid)
    cons, w2 = constrained_report(df, None if a.constrained_w_keys < 0 else a.constrained_w_keys, a.agent_name)
    est = noise_estimator_report(df)

    warns = w1 + w2
    notes = [f"Input file: {a.csv}",
             f"Objectives: T = {a.t_col}, K = {a.k_col} (both minimised)",
             "Normalisation: T^=(T-Tmin)/(Tmax-Tmin), K^=K/Kmax; bounds from --bounds file, else manuscript "
             "Table 5.2, else empirical (see bounds_source column).",
             f"noise_used_measured = {a.initial_budget:g} - Remaining_noise_budget (assumed SEAL initial budget).",
             "Constrained metrics use runs with w_keys = " + ("all values" if a.constrained_w_keys < 0
                                                              else f"{a.constrained_w_keys:g}") + "."]
    notes += [f"WARNING: {w}" for w in warns]
    sheets = {"MORL_metrics": morl, "Compiler_metrics": compiler, "Constrained_metrics": cons,
              "Noise_estimator": est, "MORL_points": points}
    out_path = os.path.join(a.out, a.xlsx_name)
    write_workbook(out_path, sheets, notes, df, points if not a.no_plots else points.iloc[0:0], df)
    if a.md:
        with open(os.path.join(a.out, "report.md"), "w") as f:
            f.write("# Evaluation report\n\n")
            for title, t in (("MORL metrics", morl), ("Compiler metrics", compiler),
                             ("Constrained agent", cons), ("Noise estimator", est)):
                f.write(f"## {title}\n\n" + md_table(t) + "\n")
    for w in warns:
        print("WARNING:", w)
    print(f"done -> {out_path}")


# --------------------------------------------------------------------------- #
def selftest():
    # single ideal point
    m, _ = morl_metrics_from_normalised([(0.0, 0.0)])
    assert (m["IC"], round(m["HV"], 6), m["SP"], round(m["EU"], 6)) == (1, 1.0, 0.0, 0.0), m
    # two extreme points (0,1) and (1,0): HV=(0.1+0.11)/1.21, SP=1, EU=0.25
    m, _ = morl_metrics_from_normalised([(0, 1), (1, 0), (1, 1), (0.5, 0.9)])
    assert m["IC"] == 3, m
    m2, _ = morl_metrics_from_normalised([(0, 1), (1, 0)])
    assert abs(m2["HV"] - 0.21 / 1.21) < 1e-9 and abs(m2["SP"] - 1.0) < 1e-9 and abs(m2["EU"] - 0.25) < 1e-3, m2
    # dominated point is removed; duplicates collapse
    assert pareto_front([(0.2, 0.2), (0.3, 0.3), (0.2, 0.2)]) == [(0.2, 0.2)]
    # estimator metrics sanity
    t = pd.DataFrame({"noise_used_estimated": [10., 20., 30.], "noise_used_measured": [11., 19., 31.]})
    e = noise_estimator_report(t).iloc[0]
    assert abs(e["MAE"] - 1.0) < 1e-9 and e["R2"] > 0.98
    # constrained: rollback bookkeeping
    d = pd.DataFrame({"benchmark": ["a", "b"], "w_ops": [1., 1.], "w_keys": [0., 0.], "noise_budget": [100., 100.],
                      "initial_ops_cost": [100., 100.], "initial_noise_estimated": [50., 150.],
                      "agent_final_ops_cost": [40., 60.], "agent_final_noise_estimated": [120., 90.],
                      "final_ops_cost": [50., 60.], "noise_used_estimated": [80., 90.],
                      "rollback_applied": [1., 0.], "noise_used_measured": [82., 91.]})
    c, _ = constrained_report(d, 0.0, "unified")
    r = c.iloc[0]
    assert r["violations"] == 1 and r["safe_violations"] == 0 and r["CR_%"] == 50.0 and r["safe_CR_%"] == 45.0, r
    assert r["within_budget"] == 1 and r["above_budget"] == 1 and r["violations_within"] == 1, r
    print("selftest OK")


if __name__ == "__main__":
    sys.exit(main())