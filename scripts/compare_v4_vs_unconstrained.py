#!/usr/bin/env python3
"""Generate COMPARISON_v4_agents_summary.xlsx — V4 agents vs unconstrained baseline."""

import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

RESULTS_DIR = os.path.join(os.path.dirname(__file__), "..", "test_results")
OUTPUT = os.path.join(RESULTS_DIR, "COMPARISON_v4_agents_summary.xlsx")

HEADER_FONT = Font(bold=True, size=11)
SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="4472C4")
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

AGENTS = {
    "v4_mb_A": {
        "label": "Margin Barrier — A [230,233,236,369,9M]",
        "budgets": [230, 233, 236, 369, 9000000],
    },
    "v4_mb_B": {
        "label": "Margin Barrier — B [230,236,369,372,9M]",
        "budgets": [230, 236, 369, 372, 9000000],
    },
    "v4_mb_film_A": {
        "label": "Margin Barrier + FiLM — A [230,233,236,369,9M]",
        "budgets": [230, 233, 236, 369, 9000000],
    },
    "v4_mb_film_B": {
        "label": "Margin Barrier + FiLM — B [230,236,369,372,9M]",
        "budgets": [230, 236, 369, 372, 9000000],
    },
    "v4_mb_film_C": {
        "label": "Margin Barrier + FiLM — C [172,230,236,369,9M]",
        "budgets": [172, 230, 236, 369, 9000000],
    },
    "v4_nmask_A": {
        "label": "Noise Masking — A [230,233,236,369,9M]",
        "budgets": [230, 233, 236, 369, 9000000],
    },
    "v4_nmask_B": {
        "label": "Noise Masking — B [230,236,369,372,9M]",
        "budgets": [230, 236, 369, 372, 9000000],
    },
    "v4_nmask_C": {
        "label": "Noise Masking — C [172,233,369,372,9M]",
        "budgets": [172, 233, 369, 372, 9000000],
    },
}

COLUMNS = [
    "Budget",
    "Num Expressions",
    "Avg Cost Reduction (%)",
    "Avg Final Noise",
    "Violations",
    "Violation Rate (%)",
    "Avg Noise Margin",
    "Avg Steps",
]

QUICK_COLUMNS = [
    "Agent",
    "Budgets",
    "Overall Avg Cost Red (%)",
    "Overall Avg Violation (%)",
    "Constrained Avg Cost Red (%)",
    "Constrained Avg Violation (%)",
    "Unconstrained Cost Red (%)",
]


def load_v4_summary(name):
    path = os.path.join(RESULTS_DIR, f"test_{name}.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["summary_per_budget"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(row)
    return rows


def load_unconstrained_noises():
    path = os.path.join(RESULTS_DIR, "test_unconstrained_ppo_baseline.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["all_results"]
    noises = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == 9000000:
            noises.append(row[6])
    return noises


def unconstrained_at_budget(noises, budget):
    n_expr = len(noises)
    violations = sum(1 for n in noises if n > budget)
    viol_rate = round(violations / n_expr * 100, 1)
    avg_noise = round(sum(noises) / n_expr, 2)
    avg_margin = round(sum(budget - n for n in noises) / n_expr, 2)
    avg_cr = 86.91
    return (budget, n_expr, avg_cr, avg_noise, violations, viol_rate, avg_margin, None)


def apply_header_style(ws, row_idx, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER


def apply_section_style(ws, row_idx, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        cell.alignment = Alignment(horizontal="left")


def apply_data_style(ws, row_idx, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")


def main():
    noises = load_unconstrained_noises()
    all_unique_budgets = sorted(set(
        b for info in AGENTS.values() for b in info["budgets"]
    ))

    wb = openpyxl.Workbook()

    # ── Sheet 1: All V4 Agents ──
    ws1 = wb.active
    ws1.title = "All V4 Agents"
    r = 1
    for i, col in enumerate(COLUMNS, 1):
        ws1.cell(row=r, column=i, value=col)
    apply_header_style(ws1, r, len(COLUMNS))
    r += 1

    # Unconstrained baseline at all unique V4 budgets
    ws1.cell(row=r, column=1, value="=== UNCONSTRAINED PPO BASELINE ===")
    apply_section_style(ws1, r, len(COLUMNS))
    r += 1
    for b in all_unique_budgets:
        row_data = unconstrained_at_budget(noises, b)
        for i, val in enumerate(row_data, 1):
            ws1.cell(row=r, column=i, value=val)
        apply_data_style(ws1, r, len(COLUMNS))
        r += 1
    r += 1  # blank separator

    # Each V4 agent
    for name, info in AGENTS.items():
        ws1.cell(row=r, column=1, value=f"=== {info['label']} ===")
        apply_section_style(ws1, r, len(COLUMNS))
        r += 1
        summary_rows = load_v4_summary(name)
        for row_data in summary_rows:
            for i, val in enumerate(row_data, 1):
                ws1.cell(row=r, column=i, value=val)
            apply_data_style(ws1, r, len(COLUMNS))
            r += 1
        r += 1  # blank separator

    for i in range(1, len(COLUMNS) + 1):
        ws1.column_dimensions[get_column_letter(i)].width = 22

    # ── Sheet 2: Quick Comparison ──
    ws2 = wb.create_sheet("Quick Comparison")
    for i, col in enumerate(QUICK_COLUMNS, 1):
        ws2.cell(row=1, column=i, value=col)
    apply_header_style(ws2, 1, len(QUICK_COLUMNS))

    quick_rows = []

    # Unconstrained baseline summary
    unc_budgets_for_summary = all_unique_budgets
    unc_viol_rates = []
    for b in unc_budgets_for_summary:
        row_data = unconstrained_at_budget(noises, b)
        unc_viol_rates.append(row_data[5])
    unc_overall_cr = 86.91
    unc_overall_viol = round(sum(unc_viol_rates) / len(unc_viol_rates), 2)
    constrained_budgets = [b for b in unc_budgets_for_summary if b < 9000000]
    unc_constr_viol_rates = []
    for b in constrained_budgets:
        row_data = unconstrained_at_budget(noises, b)
        unc_constr_viol_rates.append(row_data[5])
    unc_constr_cr = 86.91
    unc_constr_viol = round(sum(unc_constr_viol_rates) / len(unc_constr_viol_rates), 2) if unc_constr_viol_rates else 0
    quick_rows.append((
        "Unconstrained PPO (baseline)",
        ", ".join(str(b) for b in unc_budgets_for_summary),
        unc_overall_cr, unc_overall_viol,
        unc_constr_cr, unc_constr_viol,
        86.91,
    ))

    # V4 agents
    for name, info in AGENTS.items():
        summary = load_v4_summary(name)
        budgets = [row[0] for row in summary]
        all_cr = [row[2] for row in summary]
        all_viol = [row[5] for row in summary]

        overall_cr = round(sum(all_cr) / len(all_cr), 2)
        overall_viol = round(sum(all_viol) / len(all_viol), 2)

        constr_rows = [(row[2], row[5]) for row in summary if row[0] < 9000000]
        if constr_rows:
            constr_cr = round(sum(c for c, _ in constr_rows) / len(constr_rows), 2)
            constr_viol = round(sum(v for _, v in constr_rows) / len(constr_rows), 2)
        else:
            constr_cr = overall_cr
            constr_viol = overall_viol

        unconstr_rows = [row for row in summary if row[0] >= 9000000]
        unconstr_cr = unconstr_rows[0][2] if unconstr_rows else overall_cr

        quick_rows.append((
            info["label"],
            ", ".join(str(b) for b in budgets),
            overall_cr, overall_viol,
            constr_cr, constr_viol,
            unconstr_cr,
        ))

    # Sort by overall violation rate ascending (best first), unconstrained last
    v4_rows = [r for r in quick_rows if r[0] != "Unconstrained PPO (baseline)"]
    v4_rows.sort(key=lambda x: (x[3], -x[2]))
    sorted_rows = v4_rows + [quick_rows[0]]

    for row_idx, row_data in enumerate(sorted_rows, 2):
        for col_idx, val in enumerate(row_data, 1):
            ws2.cell(row=row_idx, column=col_idx, value=val)
        apply_data_style(ws2, row_idx, len(QUICK_COLUMNS))

    # Highlight best V4 agent (lowest violation, highest CR)
    if v4_rows:
        best_row_idx = 2
        for c in range(1, len(QUICK_COLUMNS) + 1):
            ws2.cell(row=best_row_idx, column=c).fill = BEST_FILL

    ws2.column_dimensions["A"].width = 50
    ws2.column_dimensions["B"].width = 35
    for i in range(3, len(QUICK_COLUMNS) + 1):
        ws2.column_dimensions[get_column_letter(i)].width = 28

    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"\nQuick Comparison (sorted by violation rate, then cost reduction):")
    print(f"{'Agent':<52} {'Budgets':<30} {'Ovr CR%':>8} {'Ovr Viol%':>10} {'Cstr CR%':>9} {'Cstr Viol%':>11} {'Unc CR%':>8}")
    print("-" * 140)
    for row in sorted_rows:
        print(f"{row[0]:<52} {row[1]:<30} {row[2]:>8.2f} {row[3]:>10.2f} {row[4]:>9.2f} {row[5]:>11.2f} {row[6]:>8.2f}")


if __name__ == "__main__":
    main()
