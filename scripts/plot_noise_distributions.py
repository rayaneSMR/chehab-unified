#!/usr/bin/env python3
"""Generate side-by-side noise distribution histograms for the manuscript.

Reads the all_results sheets from the original and threshold benchmark
Excel files and produces a PDF figure showing why the original benchmarks
are "easy" (most noise far below budgets) while the threshold benchmarks
are "hard" (noise clustered near budget boundaries).
"""

import os
import sys

import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

ORIGINAL_FILE = "test_results/original_unconstrained.xlsx"
THRESHOLD_FILE = "test_results/threshold_unconstrained.xlsx"
OUTPUT_PDF = "manuscript writing/figures-to-use/chap5/noise_distribution_comparison.pdf"
OUTPUT_PNG = "manuscript writing/figures-to-use/chap5/noise_distribution_comparison.png"
OVERLEAF_PNG = "manuscript writing/MounirMelzi_IlyesArabet___PFE_manuscript/Figures/chap5/noise_distribution_comparison.png"

BUDGETS = [230, 236, 369]
BUDGET_COLORS = {230: "#e74c3c", 236: "#e67e22", 369: "#2ecc71"}


def extract_unique_noises(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["all_results"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    noise_col = next(c for c, h in enumerate(headers, 1) if h and "Initial Noise" in str(h))
    seen = set()
    noises = []
    for r in range(2, ws.max_row + 1):
        expr_num = ws.cell(r, 2).value
        if expr_num in seen:
            continue
        seen.add(expr_num)
        n = ws.cell(r, noise_col).value
        if n is not None:
            noises.append(float(n))
    wb.close()
    return noises


def main():
    orig_noises = extract_unique_noises(ORIGINAL_FILE)
    thresh_noises = extract_unique_noises(THRESHOLD_FILE)

    os.makedirs(os.path.dirname(OUTPUT_PDF), exist_ok=True)

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(10, 4), sharey=False)

    bins_orig = np.arange(0, 500, 25)
    bins_thresh = np.arange(175, 400, 25)

    ax1.hist(orig_noises, bins=bins_orig, color="#3498db", edgecolor="white",
             alpha=0.85, zorder=3)
    for b in BUDGETS:
        ax1.axvline(b, color=BUDGET_COLORS[b], linestyle="--", linewidth=1.5,
                     label=f"B={b}", zorder=4)
    ax1.set_xlabel("Initial Noise (bits)", fontsize=11)
    ax1.set_ylabel("Number of Expressions", fontsize=11)
    ax1.set_title(f"Original Benchmarks (n={len(orig_noises)})", fontsize=12)
    ax1.legend(fontsize=9)
    ax1.grid(axis="y", alpha=0.3, zorder=0)

    ax2.hist(thresh_noises, bins=bins_thresh, color="#e74c3c", edgecolor="white",
             alpha=0.85, zorder=3)
    for b in BUDGETS:
        ax2.axvline(b, color=BUDGET_COLORS[b], linestyle="--", linewidth=1.5,
                     label=f"B={b}", zorder=4)
    ax2.set_xlabel("Initial Noise (bits)", fontsize=11)
    ax2.set_ylabel("Number of Expressions", fontsize=11)
    ax2.set_title(f"Threshold Benchmarks (n={len(thresh_noises)})", fontsize=12)
    ax2.legend(fontsize=9)
    ax2.grid(axis="y", alpha=0.3, zorder=0)

    fig.tight_layout()
    fig.savefig(OUTPUT_PDF, bbox_inches="tight", dpi=300)
    fig.savefig(OUTPUT_PNG, bbox_inches="tight", dpi=300)
    os.makedirs(os.path.dirname(OVERLEAF_PNG), exist_ok=True)
    fig.savefig(OVERLEAF_PNG, bbox_inches="tight", dpi=300)
    print(f"Saved: {OUTPUT_PDF}")
    print(f"Saved: {OUTPUT_PNG}")
    print(f"Saved: {OVERLEAF_PNG}")
    plt.close(fig)


if __name__ == "__main__":
    main()
